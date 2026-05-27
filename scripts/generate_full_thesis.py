from __future__ import annotations

import json
import math
from pathlib import Path
from statistics import mean
from typing import Dict, List

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
DOCS_DIR = ROOT / "docs"
FIGURES_DIR = DOCS_DIR / "thesis_figures"
OUTPUT_DOCX = DOCS_DIR / "SignSpeak_Thesis_6_Chapters.docx"


def load_json(path: Path) -> dict:
    with path.open("r", encoding="utf-8") as handle:
        return json.load(handle)


def parse_classification_report(path: Path) -> Dict[str, dict]:
    report: Dict[str, dict] = {}
    with path.open("r", encoding="utf-8") as handle:
        for raw_line in handle:
            line = raw_line.rstrip()
            if not line.strip():
                continue
            if line.lstrip().startswith(("precision", "accuracy", "macro avg", "weighted avg")):
                continue
            if line.startswith(" "):
                parts = line.split()
                if len(parts) == 5:
                    label = parts[0]
                    report[label] = {
                        "precision": float(parts[1]),
                        "recall": float(parts[2]),
                        "f1": float(parts[3]),
                        "support": int(parts[4]),
                    }
    return report


def set_cell_text(cell, text: str, bold: bool = False):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    run = paragraph.add_run(text)
    run.bold = bold
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if bold else WD_ALIGN_PARAGRAPH.LEFT
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def format_percent(value: float) -> str:
    return f"{value * 100:.2f}%"


def add_page_number(paragraph):
    run = paragraph.add_run()
    fld_char_begin = OxmlElement("w:fldChar")
    fld_char_begin.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = "PAGE"
    fld_char_end = OxmlElement("w:fldChar")
    fld_char_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char_begin)
    run._r.append(instr_text)
    run._r.append(fld_char_end)


def add_figure(document: Document, image_path: Path, caption: str, width: float = 6.1):
    document.add_picture(str(image_path), width=Inches(width))
    last_paragraph = document.paragraphs[-1]
    last_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption_paragraph = document.add_paragraph()
    caption_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = caption_paragraph.add_run(caption)
    run.italic = True


def add_table(document: Document, headers: List[str], rows: List[List[str]]):
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for idx, header in enumerate(headers):
        set_cell_text(table.rows[0].cells[idx], header, bold=True)
    for row_values in rows:
        row = table.add_row()
        for idx, value in enumerate(row_values):
            set_cell_text(row.cells[idx], value)
    document.add_paragraph()


def build_document() -> Path:
    word_v12 = load_json(ROOT / "artifacts" / "word_landmarks_v12" / "metrics.json")
    word_v13 = load_json(ROOT / "artifacts" / "word_landmarks_v13" / "metrics.json")
    word_v14 = load_json(ROOT / "artifacts" / "word_landmarks_v14_ft" / "metrics.json")
    alphabet_v8 = load_json(ROOT / "artifacts" / "alphabet_landmarks_v8" / "metrics.json")
    dataset_v5 = load_json(ROOT / "data" / "word_landmarks_v5_candidate" / "records.json")
    dataset_v6 = load_json(ROOT / "data" / "word_landmarks_v6_candidate" / "records.json")
    word_report = parse_classification_report(ROOT / "artifacts" / "word_landmarks_v14_ft" / "classification_report.txt")
    alphabet_report = parse_classification_report(ROOT / "artifacts" / "alphabet_landmarks_v8" / "classification_report.txt")

    weak_word_labels = sorted(word_report.items(), key=lambda item: item[1]["recall"])[:8]
    strong_word_labels = sorted(word_report.items(), key=lambda item: item[1]["recall"], reverse=True)[:8]
    lowest_alphabet = sorted(alphabet_report.items(), key=lambda item: item[1]["recall"])[:5]

    document = Document()
    styles = document.styles
    styles["Normal"].font.name = "Times New Roman"
    styles["Normal"].font.size = Pt(12)
    for style_name in ("Title", "Heading 1", "Heading 2", "Heading 3"):
        styles[style_name].font.name = "Times New Roman"

    section = document.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1.2)
    section.right_margin = Inches(1)

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title.add_run("SignSpeak AI\nA Six-Chapter Thesis on Sign Language Recognition, Learning, and Deployment")
    title_run.bold = True
    title_run.font.size = Pt(20)

    document.add_paragraph()
    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.add_run(
        "Prepared from the implemented system, trained checkpoints, and thesis model analysis artifacts"
    ).italic = True

    document.add_paragraph()
    meta = document.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta.add_run("Project: SignSpeak AI\nDate: May 17, 2026\nEnvironment: FastAPI, React, PyTorch, PostgreSQL")

    document.add_page_break()

    document.add_heading("Abstract", level=1)
    document.add_paragraph(
        "This thesis presents SignSpeak AI, a practical sign-language platform that combines isolated-word recognition, "
        "alphabet practice, live webcam inference, role-based learning workflows, and administrative oversight in a single "
        "deployable system. The project focuses on two landmark-based recognition pipelines. The first is a word-level "
        "recognizer trained on 80 American Sign Language vocabulary items represented as 32-frame MediaPipe landmark "
        "sequences with 167 features per frame. The second is an alphabet classifier trained on 26 letter classes for guided "
        "practice and live feedback. The implemented system replaces earlier RGB-heavy experimentation with lighter landmark "
        "models that are better aligned with real-time educational use."
    )
    document.add_paragraph(
        "The latest deployed word model, word_landmarks_v14_ft, was fine-tuned from v13 after expanding the candidate word "
        f"dataset from {dataset_v5['processed_videos']} to {dataset_v6['processed_videos']} processed sequences. It reached "
        f"{format_percent(word_v14['best_val_accuracy'])} best validation accuracy and {format_percent(word_v14['test_accuracy'])} "
        "held-out test accuracy across 735 test sequences. The alphabet benchmark, alphabet_landmarks_v8, achieved "
        f"{format_percent(alphabet_v8['test_accuracy'])} test accuracy on 2,183 test samples. The thesis model analysis "
        "integrates checkpoint comparisons, training curves, per-class accuracy, weak-class improvement tracking, and dataset "
        "growth visualizations to show how targeted data expansion and warm-start fine-tuning improved the final system."
    )
    document.add_paragraph(
        "Beyond recognition models, the thesis documents the end-to-end application stack, including the student, teacher, "
        "and admin experiences; PostgreSQL-backed authentication; booking and profile management; and the new admin workspace "
        "for monitoring users, teachers, classes, and active model deployment. The result is a complete software artifact that "
        "supports both academic analysis and practical operation."
    )

    document.add_page_break()
    document.add_heading("Table of Contents", level=1)
    document.add_paragraph("1. Chapter One: Introduction")
    document.add_paragraph("2. Chapter Two: Background and Related Context")
    document.add_paragraph("3. Chapter Three: System Analysis and Design")
    document.add_paragraph("4. Chapter Four: Data Pipeline and Model Development")
    document.add_paragraph("5. Chapter Five: Experimental Results and Thesis Model Analysis")
    document.add_paragraph("6. Chapter Six: Conclusion and Future Work")
    document.add_paragraph("References")

    document.add_page_break()
    document.add_heading("Chapter One: Introduction", level=1)
    document.add_paragraph(
        "SignSpeak AI was developed to address a practical gap between academic sign-language recognition experiments and "
        "deployable educational software. Many sign-language prototypes end at offline model evaluation, while many teaching "
        "applications stop at static content presentation without integrated recognition. This project bridges those concerns "
        "by combining machine-learning models, live inference services, and a role-based web platform in one codebase."
    )
    document.add_paragraph(
        "The project supports three main educational actors. Students practice alphabet handshapes, receive live predictions, "
        "book sessions, and track progress. Teachers manage profiles, classes, and learning interactions. Administrators "
        "monitor platform activity, inspect user and teacher records, review operational metrics, and supervise the active "
        "model deployment. This broad system scope is important because recognition quality alone does not determine product "
        "value; the surrounding workflow determines whether a trained model can be used consistently."
    )
    document.add_paragraph(
        "The central research objective of the thesis is to evaluate whether a landmark-based design can provide strong "
        "accuracy for isolated sign recognition while remaining efficient enough for a real-time educational platform. "
        "A secondary objective is to show how iterative dataset expansion, targeted remediation of weak classes, and "
        "fine-tuning from a prior checkpoint affect measurable model performance."
    )
    document.add_paragraph("The thesis is guided by the following research questions:")
    document.add_paragraph(
        "1. Can MediaPipe landmark sequences provide a practical foundation for both word-level recognition and alphabet practice?"
    )
    document.add_paragraph(
        "2. How much improvement is obtained by expanding the word dataset and warm-starting a new model from a stronger checkpoint?"
    )
    document.add_paragraph(
        "3. How can model quality be integrated into a complete learning platform that includes authentication, profiles, classes, and administration?"
    )
    document.add_paragraph(
        "4. Which word classes remain difficult even after substantial improvement, and what does that imply for future work?"
    )
    document.add_paragraph(
        "The scope of the thesis is limited to isolated recognition rather than continuous sign-to-sentence translation. "
        "This constraint is deliberate. The isolated setting matches the available dataset scale, the platform objectives, "
        "and the latency requirements of webcam-based educational interaction."
    )

    document.add_page_break()
    document.add_heading("Chapter Two: Background and Related Context", level=1)
    document.add_paragraph(
        "Sign-language recognition systems often balance three competing concerns: representation quality, data requirements, "
        "and inference speed. RGB video models can capture appearance-rich cues but are typically heavier, more sensitive to "
        "background conditions, and harder to align with fast interactive deployments. Landmark-based models replace much of "
        "the raw visual complexity with explicit geometric structure, making them attractive for hand-centric recognition tasks."
    )
    document.add_paragraph(
        "The SignSpeak AI project evolved through this trade-off. Early repository work included RGB-based recognition "
        "experiments, but the final platform emphasized MediaPipe landmarks because they reduce spatial redundancy and expose "
        "the motion patterns that dominate handshape-based classification. This shift also simplified live inference because "
        "the deployed backend could operate on structured hand trajectories rather than full image tensors."
    )
    document.add_paragraph(
        "Educational context also shapes the system design. Alphabet practice is not identical to word recognition. Alphabet "
        "letters are often static or short motion gestures with frequent handshape similarities, while isolated words can "
        "depend more strongly on temporal dynamics, movement direction, and transitional cues. As a result, the platform uses "
        "two specialized models rather than forcing a single architecture across all tasks."
    )
    document.add_paragraph(
        "The broader software context matters as well. A machine-learning model embedded inside a real product must coexist "
        "with authentication, database storage, profile editing, session management, and monitoring. This thesis therefore "
        "treats SignSpeak AI as a complete applied system instead of a narrow benchmark submission."
    )
    document.add_paragraph("Key platform technologies include the following:")
    document.add_paragraph("1. PyTorch for model definition, optimization, checkpointing, and evaluation.")
    document.add_paragraph("2. MediaPipe landmarks for hand tracking and geometric sequence extraction.")
    document.add_paragraph("3. FastAPI for backend inference endpoints, authentication, and admin APIs.")
    document.add_paragraph("4. React for the client application used by students, teachers, and administrators.")
    document.add_paragraph("5. PostgreSQL for persistent account and application data.")

    document.add_page_break()
    document.add_heading("Chapter Three: System Analysis and Design", level=1)
    document.add_paragraph(
        "SignSpeak AI is organized as an integrated full-stack application. The backend exposes APIs for authentication, "
        "alphabet prediction, word prediction, teacher workflows, student workflows, and admin monitoring. The frontend "
        "consumes these APIs through a role-aware application shell. The data layer persists users, profiles, schedules, and "
        "related operational records. This architecture allows machine-learning outputs to be embedded inside structured user flows."
    )
    document.add_paragraph(
        "The student experience focuses on learning interaction. Students can log in, access alphabet practice, monitor live "
        "predictions, and participate in booking or teacher-driven learning flows. The teacher experience exposes profile "
        "information, classes, and instructional presence. The newly added admin experience provides operational visibility into "
        "system-wide activity and model deployment details, including active checkpoints, dataset counts, and performance summaries."
    )
    document.add_paragraph(
        "The admin workspace is especially relevant to model lifecycle management. By separating the admin dashboard from "
        "user and teacher navigation, the platform now supports direct review of users, teacher profiles, and classes through "
        "dedicated routes. This design reduces dashboard clutter and provides clearer ownership over operational tasks."
    )
    document.add_paragraph("The system architecture can be summarized in four layers:")
    document.add_paragraph("1. Capture layer: webcam or uploaded video provides raw sign input.")
    document.add_paragraph("2. Representation layer: MediaPipe extracts normalized landmark sequences.")
    document.add_paragraph("3. Inference layer: specialized PyTorch models classify words or alphabet letters.")
    document.add_paragraph("4. Application layer: FastAPI and React present results inside learning and admin workflows.")
    document.add_paragraph(
        "For deployment, the current word service points to the checkpoint artifacts/word_landmarks_v14_ft/best.pt. "
        "Inference also uses mirrored test-time augmentation and sequence densification for short landmark dropouts, "
        "which improves runtime robustness without requiring retraining at prediction time."
    )

    document.add_page_break()
    document.add_heading("Chapter Four: Data Pipeline and Model Development", level=1)
    document.add_paragraph(
        "The final recognition stack is landmark-first. Each word sample is converted into a fixed-length sequence of 32 frames. "
        "Each frame contains 167 features that encode tracked hand information. The processed candidate dataset grew from "
        f"{dataset_v5['processed_videos']} sequences in the v5 candidate build to {dataset_v6['processed_videos']} sequences in "
        "the v6 candidate build, while the vocabulary remained fixed at 80 target words. This expansion specifically targeted "
        "classes that were weak in earlier checkpoints."
    )
    add_figure(
        document,
        FIGURES_DIR / "dataset_growth_word_clips.png",
        "Figure 1. Growth of the processed word-landmark dataset across development stages.",
    )
    document.add_paragraph(
        "The word recognizer is implemented as a bidirectional recurrent model with temporal attention. Input sequences are first "
        "normalized and projected into a higher-dimensional embedding space. A multi-layer bidirectional LSTM models temporal "
        "dependencies, and an additive attention module selects the most informative frames before final classification. "
        "The deployed v14_ft configuration uses a hidden dimension of 576, two recurrent layers, dropout of 0.4, and an "
        "attention dimension of 144."
    )
    document.add_paragraph(
        "Training uses AdamW optimization, cosine annealing of the learning rate, label smoothing, class weighting, and grouped "
        "data splitting to reduce leakage. The training pipeline was improved to fill dropped landmark frames, apply temporal "
        "cropping, and support warm-start initialization from a prior checkpoint. The v14_ft model was fine-tuned from v13 using "
        "the expanded candidate set rather than trained blindly from scratch, because the scratch run was less stable."
    )
    document.add_paragraph(
        "The alphabet model serves a different purpose and uses a lighter architecture. Instead of a recurrent attention pipeline, "
        "it uses a feed-forward multilayer perceptron over landmark-derived features, optimized for fast letter practice and live "
        "feedback. This separation of architectures reflects the difference between short static alphabet gestures and more "
        "temporally expressive word gestures."
    )
    add_table(
        document,
        ["Subsystem", "Classes", "Representation", "Model", "Best Checkpoint"],
        [
            ["Word recognition", "80 words", "32-frame landmark sequences", "BiLSTM with temporal attention", "word_landmarks_v14_ft"],
            ["Alphabet practice", "26 letters", "landmark features", "MLP classifier", "alphabet_landmarks_v8"],
        ],
    )

    document.add_page_break()
    document.add_heading("Chapter Five: Experimental Results and Thesis Model Analysis", level=1)
    document.add_paragraph(
        "This chapter presents the central thesis model analysis. The analysis is based on saved checkpoint artifacts, "
        "classification reports, and matplotlib visualizations generated directly from the repository. The focus is on three "
        "word checkpoints, v12, v13, and v14_ft, as well as the final alphabet benchmark."
    )
    add_table(
        document,
        ["Checkpoint", "Epochs", "Best Validation Accuracy", "Test Accuracy", "Test Loss"],
        [
            ["word_landmarks_v12", str(len(word_v12["history"])), format_percent(word_v12["best_val_accuracy"]), format_percent(word_v12["test_accuracy"]), f"{word_v12['test_loss']:.4f}"],
            ["word_landmarks_v13", str(len(word_v13["history"])), format_percent(word_v13["best_val_accuracy"]), format_percent(word_v13["test_accuracy"]), f"{word_v13['test_loss']:.4f}"],
            ["word_landmarks_v14_ft", str(len(word_v14["history"])), format_percent(word_v14["best_val_accuracy"]), format_percent(word_v14["test_accuracy"]), f"{word_v14['test_loss']:.4f}"],
        ],
    )
    document.add_paragraph(
        "The progression from v12 to v14_ft shows the strongest quantitative result in the project. Test accuracy improved from "
        f"{format_percent(word_v12['test_accuracy'])} in v12 to {format_percent(word_v13['test_accuracy'])} in v13 and then to "
        f"{format_percent(word_v14['test_accuracy'])} in v14_ft. The improvement from v12 to v14_ft is "
        f"{(word_v14['test_accuracy'] - word_v12['test_accuracy']) * 100:.2f} percentage points. This gain was achieved without "
        "changing the vocabulary size, which means the increase came from better data coverage and a stronger optimization path."
    )
    add_figure(
        document,
        FIGURES_DIR / "word_checkpoint_accuracy_comparison.png",
        "Figure 2. Comparison of held-out word-model accuracy across v12, v13, and v14_ft.",
    )
    add_figure(
        document,
        FIGURES_DIR / "word_checkpoint_test_loss.png",
        "Figure 3. Comparison of held-out word-model test loss across checkpoints.",
    )
    document.add_paragraph(
        "Training-curve analysis also supports the final model choice. The v14_ft history shows rapid convergence from a strong "
        "starting point, indicating that warm-start fine-tuning used prior knowledge efficiently. The best validation accuracy "
        f"peaked at {format_percent(word_v14['best_val_accuracy'])}, while the final held-out test accuracy remained closely aligned "
        "with validation behavior, suggesting good generalization rather than severe overfitting."
    )
    add_figure(
        document,
        FIGURES_DIR / "word_training_curves.png",
        "Figure 4. Training and validation behavior for the final word-model experiments.",
    )
    document.add_paragraph(
        "Per-class analysis is especially important because aggregate accuracy can hide weak categories. The final v14_ft report "
        "shows that all 80 classes exceeded 50% recall on the held-out test set, with many classes reaching perfect recall. "
        "This is a meaningful milestone because earlier checkpoints contained classes with severe failure cases."
    )
    add_table(
        document,
        ["Highest-Recall Word Classes", "Recall", "Lowest-Recall Word Classes", "Recall"],
        [
            [strong_word_labels[i][0], f"{strong_word_labels[i][1]['recall'] * 100:.0f}%", weak_word_labels[i][0], f"{weak_word_labels[i][1]['recall'] * 100:.0f}%"]
            for i in range(min(len(strong_word_labels), len(weak_word_labels)))
        ],
    )
    add_figure(
        document,
        FIGURES_DIR / "word_top_bottom_classes_v14.png",
        "Figure 5. Top and bottom per-class recall for the final word model.",
    )
    document.add_paragraph(
        "The weakest final word class was READ at 57% recall, followed by FLASHLIGHT at 67%, while classes such as BABY, BAD, "
        "CLOUD, FAMILY, LIVE, NAME, NO, and YOU reached perfect recall within the test split. The remaining difficult classes "
        "appear to be visually or temporally subtle gestures that may require additional motion diversity or better contextual cues."
    )
    document.add_paragraph(
        "Targeted weak-class remediation was successful. Earlier checkpoints struggled with labels such as ANIMAL, HOLD, LIVE, "
        "and PILLOW. After expanding the dataset and fine-tuning from v13, these classes improved substantially. This supports "
        "the practical strategy of error-driven data collection rather than indiscriminate scaling."
    )
    add_figure(
        document,
        FIGURES_DIR / "weak_class_improvement_word_models.png",
        "Figure 6. Improvement of selected weak word classes across checkpoint generations.",
    )
    document.add_paragraph(
        "The alphabet benchmark remained strong throughout the later project stages. The final alphabet_landmarks_v8 checkpoint "
        f"reached {format_percent(alphabet_v8['test_accuracy'])} test accuracy over 2,183 samples, confirming that the educational "
        "practice component is accurate enough for live guided use."
    )
    add_table(
        document,
        ["Alphabet Checkpoint", "Classes", "Best Validation Accuracy", "Test Accuracy", "Test Loss"],
        [[
            "alphabet_landmarks_v8",
            str(len(alphabet_v8["labels"])),
            format_percent(alphabet_v8["best_val_accuracy"]),
            format_percent(alphabet_v8["test_accuracy"]),
            f"{alphabet_v8['test_loss']:.4f}",
        ]],
    )
    add_figure(
        document,
        FIGURES_DIR / "alphabet_benchmark_summary.png",
        "Figure 7. Summary of the final alphabet benchmark.",
    )
    add_table(
        document,
        ["Most Challenging Alphabet Letters", "Recall"],
        [[label, f"{stats['recall'] * 100:.0f}%"] for label, stats in lowest_alphabet],
    )
    document.add_paragraph(
        "The most difficult alphabet letters in the final benchmark were H, D, N, M, and E, which is consistent with the known "
        "challenge of distinguishing similar handshapes. Even so, the lowest recall among these classes remained high enough for "
        "practical tutoring use."
    )
    document.add_paragraph(
        "Overall, the thesis model analysis demonstrates three key findings. First, landmark-based modeling is sufficient for "
        "strong educational sign-recognition performance in this scope. Second, targeted data acquisition combined with warm-start "
        "fine-tuning produced a clear measurable benefit. Third, performance reporting must include class-level analysis rather "
        "than only global accuracy, because model quality is unevenly distributed across vocabulary."
    )

    document.add_page_break()
    document.add_heading("Chapter Six: Conclusion and Future Work", level=1)
    document.add_paragraph(
        "SignSpeak AI demonstrates that a landmark-first recognition pipeline can support a complete educational sign-language "
        "application. The final platform is not only a set of trained models but also a software system with authentication, role "
        "management, live webcam services, admin tooling, and reproducible model analysis. The strongest word checkpoint reached "
        f"{format_percent(word_v14['test_accuracy'])} test accuracy over 80 classes, while the alphabet benchmark reached "
        f"{format_percent(alphabet_v8['test_accuracy'])}. These results are strong enough to justify the platform as a working "
        "applied system rather than a purely exploratory prototype."
    )
    document.add_paragraph(
        "The most important engineering lesson from the project is that architecture changes alone did not produce the final gain. "
        "The major improvement came from combining better data coverage, weak-class targeting, runtime robustness improvements, "
        "and warm-start fine-tuning. This reinforces a core applied-machine-learning principle: performance gains often come from "
        "tight iteration across data, training, and deployment rather than from a single isolated innovation."
    )
    document.add_paragraph("Several future directions are justified by the current results:")
    document.add_paragraph("1. Expand the word vocabulary beyond 80 labels with balanced support for each new class.")
    document.add_paragraph("2. Add more difficult motion-rich or semantically similar signs to stress-test the model.")
    document.add_paragraph("3. Introduce confusion-matrix driven review tools in the admin workspace for error triage.")
    document.add_paragraph("4. Explore sentence-level or continuous-sign modeling once larger aligned corpora are available.")
    document.add_paragraph("5. Evaluate cross-signer generalization under more diverse lighting, camera angle, and background conditions.")
    document.add_paragraph(
        "In conclusion, the thesis shows that SignSpeak AI is a credible end-to-end platform for isolated sign recognition and "
        "guided alphabet learning. The final system is academically analyzable, technically reproducible, and practically usable."
    )

    document.add_page_break()
    document.add_heading("References", level=1)
    document.add_paragraph("1. Project source code and training artifacts in the SignSpeak AI repository.")
    document.add_paragraph("2. Metrics and classification reports from word_landmarks_v12, word_landmarks_v13, word_landmarks_v14_ft, and alphabet_landmarks_v8.")
    document.add_paragraph("3. Dataset summaries from data/word_landmarks_v5_candidate/records.json and data/word_landmarks_v6_candidate/records.json.")
    document.add_paragraph("4. Thesis analysis notebook and exported figures in docs/thesis_model_performance_analysis.ipynb and docs/thesis_figures/.")
    document.add_paragraph("5. FastAPI, React, PyTorch, and MediaPipe software documentation used during implementation.")

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_page_number(footer)

    OUTPUT_DOCX.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT_DOCX)
    return OUTPUT_DOCX


if __name__ == "__main__":
    path = build_document()
    print(path)
